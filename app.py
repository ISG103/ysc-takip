from flask import Flask, render_template_string, request, redirect, session, url_for, send_file
import sqlite3
import os
import io
import re
from datetime import datetime, timedelta
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment

app = Flask(__name__)
app.secret_key = "super_gizli_anahtar_ysc_guvenlik_2026"
DB_NAME = "envanter.db"
UPLOAD_FOLDER = 'static/uploads'
os.makedirs(UPLOAD_FOLDER, exist_ok=True)

# GÜVENLİK ŞİFRELERİ
YONETICI_SIFRESI = "1234"
KONTROL_SIFRESI  = "1234"

# LOGO
LOGO_SRC = "https://i.hizliresim.com/zqgnfioa.png"

def veritabanini_hazirla():
    conn = sqlite3.connect(DB_NAME)
    cursor = conn.cursor()
    cursor.execute('''
        CREATE TABLE IF NOT EXISTS tupler (
            sira INTEGER,
            kod TEXT PRIMARY KEY,
            tip TEXT,
            lokasyon TEXT,
            son_kontrol TEXT,
            sonraki_kontrol TEXT,
            kontrol_eden TEXT,
            durum TEXT,
            foto_yol TEXT
        )
    ''')
    
    # Yeni eklenen Sütunlar için Veritabanı Güncellemesi (Eğer yoksa otomatik ekler)
    cursor.execute("PRAGMA table_info(tupler)")
    sutunlar = [column[1] for column in cursor.fetchall()]
    if 'dolum_tarihi' not in sutunlar:
        cursor.execute("ALTER TABLE tupler ADD COLUMN dolum_tarihi TEXT DEFAULT '-'")
    if 'periyodik_bakim' not in sutunlar:
        cursor.execute("ALTER TABLE tupler ADD COLUMN periyodik_bakim TEXT DEFAULT '-'")
        
    conn.commit()
    conn.close()

veritabanini_hazirla()

@app.after_request
def onbellegi_kapat(response):
    response.headers["Cache-Control"] = "no-cache, no-store, must-revalidate, max-age=0"
    response.headers["Pragma"] = "no-cache"
    response.headers["Expires"] = "0"
    return response

def tup_bul(gelen_kod):
    conn = sqlite3.connect(DB_NAME)
    c = conn.cursor()
    
    bugun = datetime.now().strftime("%Y-%m-%d")
    c.execute("UPDATE tupler SET durum = 'Gecikmis' WHERE sonraki_kontrol < ? AND durum != 'Gecikmis'", (bugun,))
    conn.commit()
    
    c.execute("SELECT * FROM tupler WHERE LOWER(kod) = LOWER(?)", (gelen_kod.strip(),))
    tup = c.fetchone()
    if tup:
        conn.close()
        return tup
        
    rakamlar = re.findall(r'\d+', gelen_kod)
    if rakamlar:
        sira_no = int(rakamlar[0])
        c.execute("SELECT * FROM tupler WHERE sira = ?", (sira_no,))
        tup = c.fetchone()
        if tup:
            conn.close()
            return tup
            
    conn.close()
    return None

# -------------------------------------------------------------
# HTML ŞABLONLARI
# -------------------------------------------------------------
MOBIL_HTML = '''
<!DOCTYPE html>
<html lang="tr">
<head>
    <meta name="viewport" content="width=device-width, initial-scale=1.0">
    <title>YSC / Yangın Dolabı Kontrol Kartı</title>
    <style>
        body { font-family: -apple-system, BlinkMacSystemFont, "Segoe UI", Roboto, sans-serif; background: #f1f5f9; padding: 15px; margin: 0; }
        .kart { background: white; border-radius: 14px; padding: 20px; box-shadow: 0 4px 12px rgba(0,0,0,0.08); max-width: 450px; margin: auto; }
        .logo-kutu { text-align: center; margin-bottom: 15px; }
        .logo-kutu img { max-height: 60px; max-width: 220px; object-fit: contain; }
        .baslik { font-size: 24px; font-weight: bold; color: #0f172a; margin-bottom: 5px; }
        .rozet { display: inline-block; padding: 5px 12px; border-radius: 20px; font-size: 13px; font-weight: bold; margin-bottom: 15px; }
        .Gecerli { background: #dcfce7; color: #166534; }
        .Gecikmis { background: #fee2e2; color: #991b1b; }
        .satir { border-bottom: 1px solid #e2e8f0; padding: 10px 0; font-size: 14px; }
        .satir strong { color: #475569; display: inline-block; width: 140px; }
        .kontrol-kutusu { margin-top: 20px; background: #f8fafc; border-radius: 8px; padding: 15px; border: 1px solid #cbd5e1; }
        .kontrol-kutusu label { display: block; margin: 10px 0; font-size: 15px; cursor: pointer; color: #1e293b; }
        .girdi { width: 100%; box-sizing: border-box; padding: 10px; border-radius: 6px; border: 1px solid #cbd5e1; margin-top: 5px; margin-bottom: 12px; font-size: 15px; }
        .buton { display: block; width: 100%; background: #2563eb; color: white; border: none; padding: 14px; border-radius: 8px; font-size: 16px; font-weight: bold; cursor: pointer; margin-top: 10px; }
        .mesaj { background: #e0f2fe; color: #0369a1; padding: 10px; border-radius: 6px; margin-bottom: 15px; font-size: 14px; text-align: center; }
        .yetki-alani { text-align: center; margin-top: 25px; padding-top: 15px; border-top: 1px dashed #cbd5e1; }
        .yetki-link { color: #64748b; font-size: 13px; text-decoration: none; }
        .tup-foto { width: 100%; max-height: 250px; object-fit: cover; border-radius: 8px; margin-top: 12px; border: 1px solid #cbd5e1; }
        .ozel-kutu { background: #eff6ff; border: 1px solid #bfdbfe; border-radius: 8px; padding: 10px; margin-top: 15px; }
    </style>
</head>
<body>
    <div class="kart">
        <div class="logo-kutu">
            <img src="{{ logo_src }}" alt="Ece Trafo Logo">
        </div>

        {% if basarili %}<div class="mesaj">✓ Muayene kaydı başarıyla güncellendi!</div>{% endif %}
        
        <span class="rozet {{ tup[7] }}">{{ '✓ AYLIK KONTROL GEÇERLİ' if tup[7] == 'Gecerli' else '⚠ AYLIK KONTROL GECİKMİŞ' }}</span>
        <div class="baslik">{{ tup[1] }}</div>
        <div style="color: #64748b; margin-bottom: 15px; font-size: 15px;">{{ tup[2] }}</div>
        
        <div class="satir"><strong>Sıra No:</strong> #{{ tup[0] }}</div>
        <div class="satir"><strong>Lokasyon:</strong> {{ tup[3] }}</div>
        
        <div class="ozel-kutu">
            <div class="satir" style="border:none; padding:5px 0;"><strong>Dolum Tarihi:</strong> <span style="color:#1d4ed8; font-weight:bold;">{{ tup[9] }}</span></div>
            <div class="satir" style="border:none; padding:5px 0;"><strong>Yıllık Bakım:</strong> <span style="color:#1d4ed8; font-weight:bold;">{{ tup[10] }}</span></div>
        </div>

        <div class="satir" style="margin-top:10px;"><strong>Son Aylık Kontrol:</strong> {{ tup[4] }}</div>
        <div class="satir"><strong>Sonraki Kontrol:</strong> {{ tup[5] }}</div>
        <div class="satir"><strong>Son Denetleyen:</strong> {{ tup[6] }}</div>

        {% if tup[8] %}
            <div style="margin-top: 12px;">
                <span style="font-size: 12px; color: #64748b; font-weight: bold;">Son Denetim Fotoğrafı:</span>
                <img src="/static/uploads/{{ tup[8] }}" class="tup-foto" alt="Fotoğraf">
            </div>
        {% endif %}

        {% if yetkili %}
            <form method="POST" action="/kontrol-kaydet/{{ tup[1] }}" enctype="multipart/form-data" class="kontrol-kutusu">
                <div style="display:flex; justify-content:space-between; align-items:center; margin-bottom:10px;">
                    <h4 style="margin:0; color:#0f172a;">Aylık Saha Denetimi:</h4>
                    <a href="/denetci-cikis/{{ tup[1] }}" style="font-size:11px; color:#ef4444; text-decoration:none;">(Yetkiyi Kapat)</a>
                </div>
                <label><input type="checkbox" required checked> Basınç / Hortum / Vana Durumu Normal</label>
                <label><input type="checkbox" required checked> Emniyet Pimi, Mühür veya Kilit Tam</label>
                <label><input type="checkbox" required checked> Cihaz / Dolap Temiz ve Sağlam</label>
                <label><input type="checkbox" required checked> Ekipmanın Önü Açık ve Ulaşılabilir</label>
                
                <label style="margin-top: 12px; font-weight: 600;">Kontrol Eden Personel:
                    <input type="text" name="personel" class="girdi" placeholder="Ad Soyad" required>
                </label>

                <label style="margin-top: 5px; font-weight: 600;">Fotoğraf Çek / Yükle (Opsiyonel):
                    <input type="file" name="foto" accept="image/*" capture="environment" class="girdi" style="padding: 6px;">
                </label>

                <button type="submit" class="buton">✓ Kontrolü Onayla ve Kaydet</button>
            </form>
        {% else %}
            <div class="yetki-alani">
                <a href="/denetci-giris/{{ tup[1] }}" class="yetki-link">🔒 Denetim Personeli Girişi</a>
            </div>
        {% endif %}
    </div>
</body>
</html>
'''

LOGIN_HTML = '''
<!DOCTYPE html>
<html lang="tr">
<head>
    <meta name="viewport" content="width=device-width, initial-scale=1.0">
    <title>Yönetici Girişi</title>
    <style>
        body { font-family: sans-serif; background: #f8fafc; display: flex; align-items: center; justify-content: center; height: 100vh; margin: 0; }
        .kutu { background: white; padding: 30px; border-radius: 12px; box-shadow: 0 4px 10px rgba(0,0,0,0.1); width: 320px; text-align: center; }
        input { width: 100%; box-sizing: border-box; padding: 12px; border: 1px solid #cbd5e1; border-radius: 6px; margin: 15px 0; font-size: 16px; }
        button { width: 100%; padding: 12px; background: #2563eb; color: white; border: none; border-radius: 6px; font-size: 16px; font-weight: bold; cursor: pointer; }
        .hata { color: #dc2626; font-size: 14px; margin-bottom: 10px; }
    </style>
</head>
<body>
    <div class="kutu">
        <img src="{{ logo_src }}" alt="Logo" style="max-height: 50px; margin-bottom: 15px;">
        <h3 style="margin-top:0;">Yönetici Girişi</h3>
        {% if hata %}<div class="hata">{{ hata }}</div>{% endif %}
        <form method="POST">
            <input type="password" name="sifre" placeholder="Yönetici Şifresi" required autofocus>
            <button type="submit">Giriş Yap</button>
        </form>
    </div>
</body>
</html>
'''

PANEL_HTML = '''
<!DOCTYPE html>
<html lang="tr">
<head>
    <meta charset="UTF-8">
    <title>Ece Trafo - Yangın Ekipmanları Paneli</title>
    <style>
        body { font-family: sans-serif; background: #f8fafc; padding: 25px; margin: 0; }
        .container { max-width: 1400px; margin: auto; }
        .ust-bar { display: flex; justify-content: space-between; align-items: center; margin-bottom: 20px; }
        .logo-ve-baslik { display: flex; align-items: center; gap: 18px; }
        .logo-ve-baslik img { max-height: 50px; max-width: 180px; object-fit: contain; }
        .aksiyonlar { display: flex; gap: 10px; align-items: center; flex-wrap: wrap; }
        .ozet-kutulari { display: flex; gap: 20px; margin-bottom: 25px; }
        .kutu { flex: 1; background: white; padding: 20px; border-radius: 10px; box-shadow: 0 2px 4px rgba(0,0,0,0.05); }
        .sayi { font-size: 28px; font-weight: bold; margin-top: 5px; }
        table { width: 100%; border-collapse: collapse; background: white; border-radius: 10px; overflow: hidden; box-shadow: 0 2px 4px rgba(0,0,0,0.05); }
        th, td { padding: 12px 10px; text-align: left; border-bottom: 1px solid #f1f5f9; font-size: 13px; }
        th { background: #f1f5f9; color: #475569; }
        .badge { padding: 4px 8px; border-radius: 6px; font-size: 12px; font-weight: bold; }
        .Gecerli { background: #dcfce7; color: #166534; }
        .Gecikmis { background: #fee2e2; color: #991b1b; }
        .btn-yeni { background: #2563eb; color: white; text-decoration: none; padding: 9px 15px; border-radius: 6px; font-size: 13px; font-weight: bold; display: inline-flex; align-items: center; gap: 6px; }
        .btn-excel-indir { background: #15803d; color: white; text-decoration: none; padding: 9px 15px; border-radius: 6px; font-size: 13px; font-weight: bold; display: inline-flex; align-items: center; gap: 6px; }
        .btn-excel-yukle { background: #0284c7; color: white; border: none; padding: 9px 15px; border-radius: 6px; font-size: 13px; font-weight: bold; cursor: pointer; display: inline-flex; align-items: center; gap: 6px; }
        .btn-duzenle { background: #3b82f6; color: white; text-decoration: none; padding: 5px 8px; border-radius: 5px; font-size: 12px; font-weight: bold; }
        .btn-sil { background: #ef4444; color: white; text-decoration: none; padding: 5px 8px; border-radius: 5px; font-size: 12px; font-weight: bold; margin-left: 5px; }
        .btn-cikis { background: #ef4444; color: white; text-decoration: none; padding: 9px 14px; border-radius: 6px; font-size: 13px; font-weight: bold; }
        .bilgi-banner { background: #dbeafe; border-left: 4px solid #2563eb; color: #1e40af; padding: 12px 16px; border-radius: 6px; margin-bottom: 20px; font-size: 13px; }
    </style>
</head>
<body>
    <div class="container">
        <div class="ust-bar">
            <div class="logo-ve-baslik">
                <img src="{{ logo_src }}" alt="Ece Trafo Logo">
                <h2 style="margin:0;">Yangın Ekipmanları (YSC / Dolap) Takip Paneli</h2>
            </div>
            <div class="aksiyonlar">
                <a href="/yeni-ekipman" class="btn-yeni">➕ Yeni Ekipman Ekle</a>
                <a href="/excel-indir" class="btn-excel-indir">📥 Güncel Excel İndir</a>
                
                <form method="POST" action="/excel-yukle" enctype="multipart/form-data" style="margin:0; display:inline-flex; align-items:center; gap:5px;">
                    <label class="btn-excel-yukle" style="margin:0;">
                        📤 Excel Yükle / Kalıcı Yap
                        <input type="file" name="excel_dosya" accept=".xlsx" style="display:none;" onchange="this.form.submit()">
                    </label>
                </form>

                <a href="/cikis" class="btn-cikis">Çıkış</a>
            </div>
        </div>

        {% if session.get('bildirim') %}
            <div class="bilgi-banner">✓ {{ session.pop('bildirim') }}</div>
        {% endif %}

        <div class="ozet-kutulari">
            <div class="kutu"><div>Toplam Ekipman</div><div class="sayi" style="color: #2563eb;">{{ toplam }}</div></div>
            <div class="kutu"><div>Geçerli / Kontrol Edilmiş</div><div class="sayi" style="color: #16a34a;">{{ gecerli }}</div></div>
            <div class="kutu"><div>Gecikmiş / Kontrol Bekleyen</div><div class="sayi" style="color: #dc2626;">{{ gecikmis }}</div></div>
        </div>

        <table>
            <thead>
                <tr>
                    <th style="width: 40px;">Sıra</th>
                    <th>Kod</th>
                    <th>Ekipman / Tür</th>
                    <th>Lokasyon</th>
                    <th style="background:#e0f2fe;">Dolum Tarihi</th>
                    <th style="background:#e0f2fe;">Yıllık Bakım</th>
                    <th>Son Kontrol</th>
                    <th>Sonraki Kontrol</th>
                    <th>Durum</th>
                    <th>İşlem</th>
                </tr>
            </thead>
            <tbody>
                {% for t in tupler %}
                <tr>
                    <td><span style="color: #64748b; font-weight: bold;">#{{ t[0] }}</span></td>
                    <td><strong><a href="/tup/{{ t[1] }}" target="_blank">{{ t[1] }}</a></strong></td>
                    <td>{{ t[2] }}</td>
                    <td><span style="color: #0369a1; font-weight: 500;">{{ t[3] }}</span></td>
                    <td style="font-weight:bold; color:#1d4ed8;">{{ t[9] }}</td>
                    <td style="font-weight:bold; color:#1d4ed8;">{{ t[10] }}</td>
                    <td>{{ t[4] }}</td>
                    <td>{{ t[5] }}</td>
                    <td><span class="badge {{ t[7] }}">{{ t[7] }}</span></td>
                    <td style="min-width:130px;">
                        <a href="/duzenle/{{ t[1] }}" class="btn-duzenle">✎ Düz.</a>
                        <a href="/sil/{{ t[1] }}" class="btn-sil" onclick="return confirm('{{ t[1] }} kodlu bu ekipmanı kalıcı olarak silmek istediğinize emin misiniz?');">🗑️ Sil</a>
                    </td>
                </tr>
                {% endfor %}
            </tbody>
        </table>
    </div>
</body>
</html>
'''

YENI_EKIPMAN_HTML = '''
<!DOCTYPE html>
<html lang="tr">
<head>
    <meta name="viewport" content="width=device-width, initial-scale=1.0">
    <title>Yeni Ekipman Ekle</title>
    <style>
        body { font-family: sans-serif; background: #f1f5f9; padding: 20px; margin: 0; }
        .kart { background: white; border-radius: 12px; padding: 25px; max-width: 550px; margin: auto; box-shadow: 0 4px 12px rgba(0,0,0,0.08); }
        .girdi { width: 100%; box-sizing: border-box; padding: 10px; border-radius: 6px; border: 1px solid #cbd5e1; margin-top: 5px; margin-bottom: 14px; font-size: 14px; }
        .kaydet-btn { background: #2563eb; color: white; border: none; padding: 13px 20px; border-radius: 6px; font-weight: bold; cursor: pointer; width: 100%; font-size: 16px; }
        .iptal-btn { display: block; text-align: center; margin-top: 12px; color: #64748b; text-decoration: none; font-size: 14px; }
        .etiket { font-size: 13px; font-weight: bold; color: #334155; }
        .baslik-alan { border-bottom: 1px solid #e2e8f0; padding-bottom: 10px; margin-bottom: 15px; }
        .kutu-grup { background: #f8fafc; padding: 10px; border-radius: 8px; border: 1px solid #e2e8f0; margin-bottom: 15px; }
    </style>
</head>
<body>
    <div class="kart">
        <div class="baslik-alan">
            <h2 style="margin:0; color:#0f172a;">➕ Yeni Ekipman Ekle</h2>
            <span style="font-size:12px; color:#64748b;">Otomatik Sıra No: #{{ sonraki_sira }}</span>
        </div>
        <form method="POST">
            <label class="etiket">Ekipman Kodu (Örn: TUP-048 veya DOLAP-049):</label>
            <input type="text" name="kod" class="girdi" placeholder="TUP-048" required autofocus>

            <label class="etiket">Ekipman Tipi / Kapasitesi / Türü:</label>
            <input type="text" name="tip" class="girdi" placeholder="Örn: 6 KG KKT (YANGIN TÜPÜ)" required>

            <label class="etiket">Lokasyon (Bölüm / Hat):</label>
            <input type="text" name="lokasyon" class="girdi" placeholder="Örn: Sevkiyat Alanı" required>
            
            <div class="kutu-grup">
                <div style="display: flex; gap: 10px;">
                    <div style="flex: 1;">
                        <label class="etiket">Dolum Tarihi:</label>
                        <input type="date" name="dolum_tarihi" class="girdi" required>
                    </div>
                    <div style="flex: 1;">
                        <label class="etiket">Yıllık Bakım Tarihi:</label>
                        <input type="date" name="periyodik_bakim" class="girdi" required>
                    </div>
                </div>
            </div>

            <label class="etiket">Kontrol Eden Personel:</label>
            <input type="text" name="kontrol_eden" class="girdi" value="Adil Çalışkan" required>

            <div style="display: flex; gap: 10px;">
                <div style="flex: 1;">
                    <label class="etiket">Son Aylık Kontrol:</label>
                    <input type="date" name="son_kontrol" class="girdi" value="{{ bugun }}" required>
                </div>
                <div style="flex: 1;">
                    <label class="etiket">Sonraki Aylık Kontrol:</label>
                    <input type="date" name="sonraki_kontrol" class="girdi" value="{{ sonraki_ay }}" required>
                </div>
            </div>

            <button type="submit" class="kaydet-btn">✓ Ekipmanı Sisteme Ekle</button>
            <a href="/panel" class="iptal-btn">← Vazgeç ve Panele Dön</a>
        </form>
    </div>
</body>
</html>
'''

DUZENLE_HTML = '''
<!DOCTYPE html>
<html lang="tr">
<head>
    <meta name="viewport" content="width=device-width, initial-scale=1.0">
    <title>Ekipman Düzenle (Yönetici)</title>
    <style>
        body { font-family: sans-serif; background: #f1f5f9; padding: 20px; margin: 0; }
        .kart { background: white; border-radius: 12px; padding: 25px; max-width: 550px; margin: auto; box-shadow: 0 4px 12px rgba(0,0,0,0.08); }
        .girdi { width: 100%; box-sizing: border-box; padding: 10px; border-radius: 6px; border: 1px solid #cbd5e1; margin-top: 5px; margin-bottom: 14px; font-size: 14px; }
        .kaydet-btn { background: #16a34a; color: white; border: none; padding: 13px 20px; border-radius: 6px; font-weight: bold; cursor: pointer; width: 100%; font-size: 16px; }
        .iptal-btn { display: block; text-align: center; margin-top: 12px; color: #64748b; text-decoration: none; font-size: 14px; }
        .etiket { font-size: 13px; font-weight: bold; color: #334155; }
        .baslik-alan { border-bottom: 1px solid #e2e8f0; padding-bottom: 10px; margin-bottom: 15px; }
        .kutu-grup { background: #f8fafc; padding: 10px; border-radius: 8px; border: 1px solid #e2e8f0; margin-bottom: 15px; }
    </style>
</head>
<body>
    <div class="kart">
        <div class="baslik-alan">
            <h2 style="margin:0; color:#0f172a;">{{ tup[1] }} Düzenleme</h2>
            <span style="font-size:12px; color:#64748b;">Sıra No: #{{ tup[0] }}</span>
        </div>
        <form method="POST">
            <label class="etiket">Ekipman Tipi / Kapasitesi / Türü:</label>
            <input type="text" name="tip" class="girdi" value="{{ tup[2] }}" required>

            <label class="etiket">Lokasyon (Bina / Bölüm / Hat):</label>
            <input type="text" name="lokasyon" class="girdi" value="{{ tup[3] }}" required>
            
            <div class="kutu-grup">
                <div style="display: flex; gap: 10px;">
                    <div style="flex: 1;">
                        <label class="etiket">Dolum Tarihi:</label>
                        <input type="date" name="dolum_tarihi" class="girdi" value="{{ tup[9] }}">
                    </div>
                    <div style="flex: 1;">
                        <label class="etiket">Yıllık Bakım Tarihi:</label>
                        <input type="date" name="periyodik_bakim" class="girdi" value="{{ tup[10] }}">
                    </div>
                </div>
            </div>

            <label class="etiket">Kontrol Eden Personel (Aylık):</label>
            <input type="text" name="kontrol_eden" class="girdi" value="{{ tup[6] }}" placeholder="Örn: Personel Adı" required>

            <div style="display: flex; gap: 10px;">
                <div style="flex: 1;">
                    <label class="etiket">Son Aylık Kontrol:</label>
                    <input type="date" name="son_kontrol" class="girdi" value="{{ tup[4] }}" required>
                </div>
                <div style="flex: 1;">
                    <label class="etiket">Sonraki Aylık Kontrol:</label>
                    <input type="date" name="sonraki_kontrol" class="girdi" value="{{ tup[5] }}" required>
                </div>
            </div>

            <label class="etiket">Genel Durum:</label>
            <select name="durum" class="girdi">
                <option value="Gecerli" {% if tup[7] == 'Gecerli' %}selected{% endif %}>✓ Geçerli (Denetimi Yapılmış)</option>
                <option value="Gecikmis" {% if tup[7] == 'Gecikmis' %}selected{% endif %}>⚠ Gecikmiş (Kontrol Bekliyor)</option>
            </select>

            <button type="submit" class="kaydet-btn">✓ Tüm Değişiklikleri Kaydet</button>
            <a href="/panel" class="iptal-btn">← Değişiklik Yapmadan Panele Dön</a>
        </form>
    </div>
</body>
</html>
'''

# -------------------------------------------------------------
# ROTALAR
# -------------------------------------------------------------
@app.route('/')
def ana_sayfa():
    return redirect('/panel')

@app.route('/tup/<kod>')
def tup_detay(kod):
    basarili = request.args.get('kaydedildi', False)
    yetkili = session.get('denetci_yetkisi', False)
    tup = tup_bul(kod)
    if not tup:
        return f"<h3>Ekipman bulunamadı! (Aranan Kod: {kod})</h3>", 404
    return render_template_string(MOBIL_HTML, tup=tup, basarili=basarili, yetkili=yetkili, logo_src=LOGO_SRC)

@app.route('/denetci-giris/<kod>', methods=['GET', 'POST'])
def denetci_giris(kod):
    tup = tup_bul(kod)
    asıl_kod = tup[1] if tup else kod
    if request.method == 'POST':
        if request.form.get('pin') == KONTROL_SIFRESI:
            session['denetci_yetkisi'] = True
            return redirect(f"/tup/{asıl_kod}")
    DENETCI_HTML = '''<div style="font-family:sans-serif; text-align:center; padding:40px;"><form method="POST"><input type="password" name="pin" placeholder="PIN Kodu" style="padding:10px; font-size:18px;"><br><br><button type="submit" style="padding:10px 20px;">Giriş</button></form></div>'''
    return render_template_string(DENETCI_HTML)

@app.route('/denetci-cikis/<kod>')
def denetci_cikis(kod):
    session.pop('denetci_yetkisi', None)
    return redirect(f"/tup/{kod}")

@app.route('/kontrol-kaydet/<kod>', methods=['POST'])
def kontrol_kaydet(kod):
    if not session.get('denetci_yetkisi'):
        return "Yetkisiz işlem!", 403
    tup = tup_bul(kod)
    if not tup:
        return "Ekipman bulunamadı!", 404
    asıl_kod = tup[1]

    personel = request.form.get('personel', 'Yetkili Personel')
    su_an = datetime.now().strftime("%Y-%m-%d %H:%M")
    sonraki = (datetime.now() + timedelta(days=30)).strftime("%Y-%m-%d")
    
    foto_adi = ""
    if 'foto' in request.files:
        foto = request.files['foto']
        if foto and foto.filename != '':
            uzanti = os.path.splitext(foto.filename)[1]
            foto_adi = f"{asıl_kod}_{int(datetime.now().timestamp())}{uzanti}"
            foto.save(os.path.join(UPLOAD_FOLDER, foto_adi))

    conn = sqlite3.connect(DB_NAME)
    c = conn.cursor()
    if foto_adi != "":
        c.execute('''UPDATE tupler SET son_kontrol = ?, sonraki_kontrol = ?, kontrol_eden = ?, durum = 'Gecerli', foto_yol = ? WHERE kod = ?''', (su_an, sonraki, personel, foto_adi, asıl_kod))
    else:
        c.execute('''UPDATE tupler SET son_kontrol = ?, sonraki_kontrol = ?, kontrol_eden = ?, durum = 'Gecerli' WHERE kod = ?''', (su_an, sonraki, personel, asıl_kod))
    conn.commit()
    conn.close()
    return redirect(f"/tup/{asıl_kod}?kaydedildi=1")

@app.route('/login', methods=['GET', 'POST'])
def login():
    hata = None
    if request.method == 'POST':
        if request.form.get('sifre') == YONETICI_SIFRESI:
            session['giris_yapti'] = True
            return redirect('/panel')
        else:
            hata = "Hatalı şifre!"
    return render_template_string(LOGIN_HTML, hata=hata, logo_src=LOGO_SRC)

@app.route('/cikis')
def cikis():
    session.pop('giris_yapti', None)
    return redirect('/login')

@app.route('/panel')
def yonetici_paneli():
    if not session.get('giris_yapti'):
        return redirect('/login')
        
    conn = sqlite3.connect(DB_NAME)
    c = conn.cursor()
    bugun = datetime.now().strftime("%Y-%m-%d")
    c.execute("UPDATE tupler SET durum = 'Gecikmis' WHERE sonraki_kontrol < ?", (bugun,))
    conn.commit()
    c.execute("SELECT * FROM tupler ORDER BY sira ASC")
    tupler = c.fetchall()
    toplam = len(tupler)
    gecerli = sum(1 for t in tupler if t[7] == 'Gecerli')
    gecikmis = toplam - gecerli
    conn.close()
    return render_template_string(PANEL_HTML, tupler=tupler, toplam=toplam, gecerli=gecerli, gecikmis=gecikmis, logo_src=LOGO_SRC)

@app.route('/yeni-ekipman', methods=['GET', 'POST'])
def yeni_ekipman():
    if not session.get('giris_yapti'):
        return redirect('/login')

    conn = sqlite3.connect(DB_NAME)
    c = conn.cursor()

    if request.method == 'POST':
        kod = request.form.get('kod').strip().upper()
        tip = request.form.get('tip').strip()
        lokasyon = request.form.get('lokasyon').strip()
        dolum_tarihi = request.form.get('dolum_tarihi')
        periyodik_bakim = request.form.get('periyodik_bakim')
        kontrol_eden = request.form.get('kontrol_eden').strip()
        son_kontrol = request.form.get('son_kontrol')
        sonraki_kontrol = request.form.get('sonraki_kontrol')

        c.execute("SELECT MAX(sira) FROM tupler")
        max_sira = c.fetchone()[0] or 0
        yeni_sira = max_sira + 1

        c.execute('''
            INSERT INTO tupler (sira, kod, tip, lokasyon, son_kontrol, sonraki_kontrol, kontrol_eden, durum, foto_yol, dolum_tarihi, periyodik_bakim)
            VALUES (?, ?, ?, ?, ?, ?, ?, 'Gecerli', '', ?, ?)
        ''', (yeni_sira, kod, tip, lokasyon, son_kontrol, sonraki_kontrol, kontrol_eden, dolum_tarihi, periyodik_bakim))
        conn.commit()
        conn.close()

        session['bildirim'] = f"{kod} kodlu yeni ekipman başarıyla sisteme eklendi!"
        return redirect('/panel')

    c.execute("SELECT MAX(sira) FROM tupler")
    max_sira = c.fetchone()[0] or 0
    sonraki_sira = max_sira + 1
    conn.close()

    bugun = datetime.now().strftime("%Y-%m-%d")
    sonraki_ay = (datetime.now() + timedelta(days=30)).strftime("%Y-%m-%d")

    return render_template_string(YENI_EKIPMAN_HTML, sonraki_sira=sonraki_sira, bugun=bugun, sonraki_ay=sonraki_ay)

@app.route('/excel-indir')
def excel_indir():
    if not session.get('giris_yapti'):
        return redirect('/login')

    conn = sqlite3.connect(DB_NAME)
    c = conn.cursor()
    c.execute("SELECT sira, kod, tip, lokasyon, son_kontrol, sonraki_kontrol, kontrol_eden, durum, dolum_tarihi, periyodik_bakim FROM tupler ORDER BY sira ASC")
    veriler = c.fetchall()
    conn.close()

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "YSC Denetim Listesi"

    ws.merge_cells('A1:J1')
    ws['A1'] = "ECE TRAFO - YANGIN EKİPMANLARI PERİYODİK KONTROL RAPORU"
    ws['A1'].font = Font(name="Arial", size=14, bold=True, color="FFFFFF")
    ws['A1'].fill = PatternFill(start_color="1E3A8A", fill_type="solid")
    ws['A1'].alignment = Alignment(horizontal="center", vertical="center")
    ws.row_dimensions[1].height = 35

    sutunlar = ["Sıra No", "Ekipman Kodu", "Tip / Tür", "Lokasyon", "Son Kontrol Tarihi", "Sonraki Kontrol", "Denetleyen", "Durum", "Dolum Tarihi", "Yıllık Bakım"]
    ws.append([])
    ws.append(sutunlar)

    for col in range(1, 11):
        hucre = ws.cell(row=3, column=col)
        hucre.font = Font(name="Arial", size=11, bold=True, color="FFFFFF")
        hucre.fill = PatternFill(start_color="3B82F6", fill_type="solid")
        hucre.alignment = Alignment(horizontal="center", vertical="center")
    ws.row_dimensions[3].height = 25

    for satir_idx, row in enumerate(veriler, start=4):
        ws.append(list(row))
        durum_hucre = ws.cell(row=satir_idx, column=8)
        if row[7] == 'Gecerli':
            durum_hucre.fill = PatternFill(start_color="DCFCE7", fill_type="solid")
            durum_hucre.font = Font(color="15803D", bold=True)
        else:
            durum_hucre.fill = PatternFill(start_color="FEE2E2", fill_type="solid")
            durum_hucre.font = Font(color="B91C1C", bold=True)

    for col in ws.columns:
        max_len = max(len(str(cell.value or '')) for cell in col)
        col_letter = openpyxl.utils.get_column_letter(col[0].column)
        ws.column_dimensions[col_letter].width = max(max_len + 4, 12)

    buffer = io.BytesIO()
    wb.save(buffer)
    buffer.seek(0)
    
    dosya_adi = f"Ece_Trafo_YSC_Raporu_{datetime.now().strftime('%Y%m%d')}.xlsx"
    return send_file(buffer, as_attachment=True, download_name=dosya_adi, mimetype='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet')

@app.route('/excel-yukle', methods=['POST'])
def excel_yukle():
    if not session.get('giris_yapti'):
        return redirect('/login')

    if 'excel_dosya' not in request.files:
        return redirect('/panel')
        
    file = request.files['excel_dosya']
    if file.filename == '':
        return redirect('/panel')

    if file:
        wb = openpyxl.load_workbook(file)
        ws = wb.active
        
        conn = sqlite3.connect(DB_NAME)
        c = conn.cursor()
        
        c.execute("DELETE FROM tupler")
        
        guncellenen = 0
        for row in ws.iter_rows(min_row=4, values_only=True):
            if not row or not row[1]:
                continue
            sira = row[0]
            kod = str(row[1]).strip()
            tip = str(row[2]).strip() if row[2] else ''
            lokasyon = str(row[3]).strip() if row[3] else ''
            son_kontrol = str(row[4])[:10] if row[4] else ''
            sonraki_kontrol = str(row[5])[:10] if row[5] else ''
            kontrol_eden = str(row[6]).strip() if row[6] else ''
            durum = str(row[7]).strip() if row[7] else 'Gecerli'
            dolum_tarihi = str(row[8])[:10] if len(row) > 8 and row[8] else '-'
            periyodik_bakim = str(row[9])[:10] if len(row) > 9 and row[9] else '-'
            
            c.execute('''
                INSERT INTO tupler (sira, kod, tip, lokasyon, son_kontrol, sonraki_kontrol, kontrol_eden, durum, foto_yol, dolum_tarihi, periyodik_bakim)
                VALUES (?, ?, ?, ?, ?, ?, ?, ?, '', ?, ?)
            ''', (sira, kod, tip, lokasyon, son_kontrol, sonraki_kontrol, kontrol_eden, durum, dolum_tarihi, periyodik_bakim))
            guncellenen += 1
            
        conn.commit()
        conn.close()
        session['bildirim'] = f"{guncellenen} adet ekipman verisi Excel'den tertemiz şekilde sisteme aktarıldı!"
        
    return redirect('/panel')

@app.route('/duzenle/<kod>', methods=['GET', 'POST'])
def duzenle(kod):
    if not session.get('giris_yapti'):
        return redirect('/login')
        
    tup = tup_bul(kod)
    if not tup:
        return "Ekipman bulunamadı!", 404
    asıl_kod = tup[1]

    conn = sqlite3.connect(DB_NAME)
    c = conn.cursor()
    if request.method == 'POST':
        yeni_tip = request.form.get('tip')
        yeni_lokasyon = request.form.get('lokasyon')
        yeni_kontrol_eden = request.form.get('kontrol_eden')
        yeni_son_kontrol = request.form.get('son_kontrol')
        yeni_sonraki_kontrol = request.form.get('sonraki_kontrol')
        yeni_durum = request.form.get('durum')
        yeni_dolum = request.form.get('dolum_tarihi')
        yeni_bakim = request.form.get('periyodik_bakim')

        c.execute('''
            UPDATE tupler 
            SET tip = ?, lokasyon = ?, kontrol_eden = ?, son_kontrol = ?, sonraki_kontrol = ?, durum = ?, dolum_tarihi = ?, periyodik_bakim = ?
            WHERE kod = ?
        ''', (yeni_tip, yeni_lokasyon, yeni_kontrol_eden, yeni_son_kontrol, yeni_sonraki_kontrol, yeni_durum, yeni_dolum, yeni_bakim, asıl_kod))
        conn.commit()
        conn.close()
        return redirect('/panel')
        
    conn.close()
    return render_template_string(DUZENLE_HTML, tup=tup)

@app.route('/sil/<kod>')
def ekipman_sil(kod):
    if not session.get('giris_yapti'):
        return redirect('/login')
        
    conn = sqlite3.connect(DB_NAME)
    c = conn.cursor()
    c.execute("DELETE FROM tupler WHERE kod = ?", (kod,))
    conn.commit()
    conn.close()
    
    session['bildirim'] = f"{kod} kodlu ekipman kalıcı olarak silindi!"
    return redirect('/panel')

if __name__ == '__main__':
    app.run(host='0.0.0.0', port=5000, debug=True)
